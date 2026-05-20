from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- COLOR PALETTE (Dark Premium Theme) ---
BG_DARK = RGBColor(0x0A, 0x0A, 0x0A)
BG_SURFACE = RGBColor(0x11, 0x11, 0x11)
BG_CARD = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT_BLUE = RGBColor(0x3B, 0x82, 0xF6)
ACCENT_PURPLE = RGBColor(0x7C, 0x3A, 0xED)
TEXT_WHITE = RGBColor(0xF8, 0xFA, 0xFC)
TEXT_MUTED = RGBColor(0x94, 0xA3, 0xB8)
TEXT_GREEN = RGBColor(0x22, 0xC5, 0x5E)
BORDER_COLOR = RGBColor(0x33, 0x33, 0x33)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def set_slide_bg(slide, color=BG_DARK):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_textbox(slide, left, top, width, height, text, font_size=18, bold=False, color=TEXT_WHITE, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return tf

def add_paragraph(tf, text, font_size=16, bold=False, color=TEXT_WHITE, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return p

def add_card(slide, left, top, width, height, fill_color=BG_CARD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = BORDER_COLOR
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape

def add_accent_line(slide, left, top, width):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.06))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()
    return shape

# ============================================================
# SLIDE 1: TITLE SLIDE
# ============================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide1)
add_accent_line(slide1, 0, 0, 13.333)

add_textbox(slide1, 1, 1.5, 11, 1.2, "AGENTIC AI-BASED", font_size=44, bold=True, color=TEXT_WHITE, alignment=PP_ALIGN.CENTER, font_name="Calibri")
add_textbox(slide1, 1, 2.5, 11, 1.2, "TRAVEL PLANNING ASSISTANT", font_size=48, bold=True, color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER, font_name="Calibri")
add_textbox(slide1, 1, 3.7, 11, 0.7, "Using LangChain, LangGraph & Google Gemini", font_size=22, bold=False, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

add_card(slide1, 3.5, 4.8, 6.3, 1.5, BG_CARD)
add_textbox(slide1, 3.7, 4.9, 5.9, 0.5, "PROJECT REPORT", font_size=14, bold=True, color=ACCENT_PURPLE, alignment=PP_ALIGN.CENTER)
add_textbox(slide1, 3.7, 5.3, 5.9, 0.4, "Domain: Travel / Tourism", font_size=16, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)
add_textbox(slide1, 3.7, 5.7, 5.9, 0.4, "Submitted to: Labmentix", font_size=16, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

add_textbox(slide1, 1, 6.8, 11, 0.4, "© 2026  |  Python • LangChain • Streamlit • SQLite • Open-Meteo API", font_size=12, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 2: TABLE OF CONTENTS
# ============================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide2)
add_accent_line(slide2, 0.8, 0.6, 2)
add_textbox(slide2, 0.8, 0.7, 5, 0.7, "TABLE OF CONTENTS", font_size=32, bold=True, color=TEXT_WHITE)

toc_items = [
    ("01", "Problem Statement & Business Use Cases"),
    ("02", "Project Objectives"),
    ("03", "Technology Stack & Architecture"),
    ("04", "Data Sources & Database Design"),
    ("05", "LangChain Tools Implementation"),
    ("06", "Agentic AI Workflow"),
    ("07", "Streamlit UI/UX Design"),
    ("08", "Expected Output & Demo"),
    ("09", "Code Quality & Best Practices"),
    ("10", "Conclusion & Future Scope"),
]

for i, (num, title) in enumerate(toc_items):
    y = 1.8 + i * 0.5
    add_card(slide2, 0.8, y, 11.5, 0.42, BG_CARD)
    add_textbox(slide2, 1.0, y + 0.02, 0.6, 0.38, num, font_size=14, bold=True, color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER)
    add_textbox(slide2, 1.7, y + 0.02, 10, 0.38, title, font_size=16, color=TEXT_WHITE)

# ============================================================
# SLIDE 3: PROBLEM STATEMENT
# ============================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide3)
add_accent_line(slide3, 0.8, 0.6, 2)
add_textbox(slide3, 0.8, 0.7, 6, 0.7, "PROBLEM STATEMENT", font_size=32, bold=True, color=TEXT_WHITE)

add_card(slide3, 0.8, 1.7, 11.5, 2.2, BG_CARD)
tf3 = add_textbox(slide3, 1.1, 1.8, 11, 0.4, "THE CHALLENGE", font_size=14, bold=True, color=ACCENT_PURPLE)
add_paragraph(tf3, "", font_size=8)
add_paragraph(tf3, "Planning a trip requires choosing flights, hotels, and attractions while considering", font_size=15, color=TEXT_MUTED)
add_paragraph(tf3, "time, budget, weather, distance, and personal preferences.", font_size=15, color=TEXT_MUTED)
add_paragraph(tf3, "", font_size=8)
add_paragraph(tf3, "Travelers often switch between multiple websites, compare inconsistent information,", font_size=15, color=TEXT_MUTED)
add_paragraph(tf3, "and manually build itineraries that may be inefficient, unrealistic, or incomplete.", font_size=15, color=TEXT_MUTED)

add_card(slide3, 0.8, 4.2, 11.5, 2.5, BG_CARD)
tf3b = add_textbox(slide3, 1.1, 4.3, 11, 0.4, "BUSINESS USE CASES", font_size=14, bold=True, color=ACCENT_PURPLE)
biz_items = [
    "✦  Reduce customer support workload for travel agencies",
    "✦  Provide hyper-personalized recommendations at scale",
    "✦  Automate itinerary design with real-time data",
    "✦  Improve customer satisfaction & save time and money",
    "✦  Adopted by: MakeMyTrip, Booking.com, ClearTrip, Ixigo",
]
for item in biz_items:
    add_paragraph(tf3b, item, font_size=15, color=TEXT_MUTED)

# ============================================================
# SLIDE 4: PROJECT OBJECTIVES
# ============================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide4)
add_accent_line(slide4, 0.8, 0.6, 2)
add_textbox(slide4, 0.8, 0.7, 6, 0.7, "PROJECT OBJECTIVES", font_size=32, bold=True, color=TEXT_WHITE)

add_card(slide4, 0.8, 1.7, 5.5, 5, BG_CARD)
tf4a = add_textbox(slide4, 1.1, 1.8, 5, 0.4, "PRIMARY OBJECTIVES", font_size=14, bold=True, color=TEXT_GREEN)
primary = [
    "Build an agentic AI system using LangChain",
    "Integrate tools: Flight, Hotel, Places, Weather",
    "Enable multi-step reasoning (ReAct Agent)",
    "Generate structured itineraries with:",
    "   • Day-wise plan",
    "   • Accommodation details",
    "   • Weather expectations",
    "   • Budget estimation",
]
for item in primary:
    add_paragraph(tf4a, item, font_size=14, color=TEXT_MUTED)

add_card(slide4, 6.8, 1.7, 5.5, 5, BG_CARD)
tf4b = add_textbox(slide4, 7.1, 1.8, 5, 0.4, "SECONDARY OBJECTIVES", font_size=14, bold=True, color=ACCENT_BLUE)
secondary = [
    "Implement filtering, ranking & optimization",
    "  (cheapest flight, highest-rated hotel)",
    "Ensure the system can justify decisions",
    "  ('Why we selected this?')",
    "Provide output in JSON + human-readable format",
    "Build a beautiful Streamlit interface",
    "Follow PEP 8 coding standards",
    "Comprehensive documentation (README)",
]
for item in secondary:
    add_paragraph(tf4b, item, font_size=14, color=TEXT_MUTED)

# ============================================================
# SLIDE 5: TECHNOLOGY STACK
# ============================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide5)
add_accent_line(slide5, 0.8, 0.6, 2)
add_textbox(slide5, 0.8, 0.7, 8, 0.7, "TECHNOLOGY STACK & ARCHITECTURE", font_size=32, bold=True, color=TEXT_WHITE)

techs = [
    ("Python 3.14", "Core programming language"),
    ("LangChain + LangGraph", "Agentic AI framework & orchestration"),
    ("Google Gemini 2.5 Flash", "Large Language Model for reasoning"),
    ("Streamlit", "Frontend web application framework"),
    ("SQLite", "Local relational database for data storage"),
    ("Open-Meteo API", "Free real-time weather forecasting"),
]

for i, (tech, desc) in enumerate(techs):
    col = i % 3
    row = i // 3
    x = 0.8 + col * 4.0
    y = 1.7 + row * 2.7
    add_card(slide5, x, y, 3.7, 2.3, BG_CARD)
    add_textbox(slide5, x + 0.2, y + 0.3, 3.3, 0.5, tech, font_size=20, bold=True, color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER)
    add_textbox(slide5, x + 0.2, y + 1.0, 3.3, 1.0, desc, font_size=14, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 6: DATA SOURCES & DATABASE
# ============================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide6)
add_accent_line(slide6, 0.8, 0.6, 2)
add_textbox(slide6, 0.8, 0.7, 8, 0.7, "DATA SOURCES & DATABASE DESIGN", font_size=32, bold=True, color=TEXT_WHITE)

add_card(slide6, 0.8, 1.7, 7.5, 5, BG_CARD)
tf6a = add_textbox(slide6, 1.1, 1.8, 7, 0.4, "JSON DATASETS → SQLite DATABASE", font_size=14, bold=True, color=ACCENT_PURPLE)
db_items = [
    "flights.json → flights table",
    "  Fields: flight_id, airline, source, destination, departure,",
    "  duration_hrs, price_inr",
    "",
    "hotels.json → hotels table",
    "  Fields: hotel_id, name, city, rating, price_per_night_inr,",
    "  amenities",
    "",
    "places.json → places table",
    "  Fields: place_id, name, city, type, rating, description",
    "",
    "SQL Optimizations:",
    "  • Indexed on (source, destination) for flights",
    "  • Indexed on (city) for hotels & places",
    "  • Normalized tables — no redundancy",
]
for item in db_items:
    add_paragraph(tf6a, item, font_size=13, color=TEXT_MUTED)

add_card(slide6, 8.8, 1.7, 3.7, 5, BG_CARD)
tf6b = add_textbox(slide6, 9.1, 1.8, 3.3, 0.4, "LIVE WEATHER API", font_size=14, bold=True, color=TEXT_GREEN)
add_paragraph(tf6b, "", font_size=6)
add_paragraph(tf6b, "Open-Meteo API", font_size=16, bold=True, color=TEXT_WHITE)
add_paragraph(tf6b, "", font_size=6)
add_paragraph(tf6b, "✦ No API key required", font_size=13, color=TEXT_MUTED)
add_paragraph(tf6b, "✦ 7-day forecast", font_size=13, color=TEXT_MUTED)
add_paragraph(tf6b, "✦ Temperature data", font_size=13, color=TEXT_MUTED)
add_paragraph(tf6b, "✦ Timezone-aware", font_size=13, color=TEXT_MUTED)
add_paragraph(tf6b, "", font_size=10)
add_paragraph(tf6b, "Endpoint:", font_size=12, bold=True, color=ACCENT_BLUE)
add_paragraph(tf6b, "api.open-meteo.com", font_size=12, color=TEXT_MUTED)
add_paragraph(tf6b, "/v1/forecast", font_size=12, color=TEXT_MUTED)

# ============================================================
# SLIDE 7: LANGCHAIN TOOLS
# ============================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide7)
add_accent_line(slide7, 0.8, 0.6, 2)
add_textbox(slide7, 0.8, 0.7, 8, 0.7, "LANGCHAIN TOOLS IMPLEMENTATION", font_size=32, bold=True, color=TEXT_WHITE)

tools_data = [
    ("✈️  FLIGHT SEARCH", "Queries SQLite for flights by source → destination. Suggests cheapest/fastest option.", ACCENT_BLUE),
    ("🏨  HOTEL RECOMMENDATION", "Filters hotels by city, ranks by rating & price. Returns top matches.", ACCENT_PURPLE),
    ("📍  PLACES DISCOVERY", "Recommends tourist attractions based on city, type & rating.", TEXT_GREEN),
    ("🌤️  WEATHER LOOKUP", "Calls Open-Meteo API with lat/long. Returns 7-day forecast.", RGBColor(0xF5, 0x9E, 0x0B)),
    ("💰  BUDGET ESTIMATION", "Sums flight + hotel + daily local expenses. Returns structured breakdown.", RGBColor(0xEF, 0x44, 0x44)),
]

for i, (title, desc, color) in enumerate(tools_data):
    y = 1.6 + i * 1.1
    add_card(slide7, 0.8, y, 11.5, 0.95, BG_CARD)
    add_textbox(slide7, 1.1, y + 0.1, 4, 0.4, title, font_size=16, bold=True, color=color)
    add_textbox(slide7, 5.5, y + 0.15, 6.5, 0.7, desc, font_size=14, color=TEXT_MUTED)

# ============================================================
# SLIDE 8: AGENTIC AI WORKFLOW
# ============================================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide8)
add_accent_line(slide8, 0.8, 0.6, 2)
add_textbox(slide8, 0.8, 0.7, 8, 0.7, "AGENTIC AI WORKFLOW", font_size=32, bold=True, color=TEXT_WHITE)

steps = [
    ("01", "USER INPUT", "User enters travel query\n(destination, dates, preferences)"),
    ("02", "AGENT REASONING", "LangGraph ReAct agent\nanalyzes the request"),
    ("03", "TOOL CALLING", "Agent autonomously selects\nand calls relevant tools"),
    ("04", "DATA RETRIEVAL", "Tools query SQLite DB\nand Open-Meteo API"),
    ("05", "ITINERARY GEN", "Agent constructs 3-7 day\nstructured itinerary"),
    ("06", "OUTPUT", "Final answer with flights,\nhotels, weather & budget"),
]

for i, (num, title, desc) in enumerate(steps):
    col = i % 3
    row = i // 3
    x = 0.8 + col * 4.1
    y = 1.7 + row * 2.8
    add_card(slide8, x, y, 3.7, 2.4, BG_CARD)
    add_textbox(slide8, x + 0.15, y + 0.15, 0.5, 0.4, num, font_size=22, bold=True, color=ACCENT_BLUE)
    add_textbox(slide8, x + 0.7, y + 0.15, 2.8, 0.4, title, font_size=16, bold=True, color=TEXT_WHITE)
    add_textbox(slide8, x + 0.2, y + 0.7, 3.3, 1.5, desc, font_size=13, color=TEXT_MUTED)

# ============================================================
# SLIDE 9: STREAMLIT UI/UX
# ============================================================
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide9)
add_accent_line(slide9, 0.8, 0.6, 2)
add_textbox(slide9, 0.8, 0.7, 8, 0.7, "STREAMLIT UI / UX DESIGN", font_size=32, bold=True, color=TEXT_WHITE)

add_card(slide9, 0.8, 1.7, 5.5, 5, BG_CARD)
tf9a = add_textbox(slide9, 1.1, 1.8, 5, 0.4, "DESIGN SYSTEM", font_size=14, bold=True, color=ACCENT_PURPLE)
ui_items = [
    "✦ Deep dark theme (#0A0A0A base)",
    "✦ Space Grotesk + Inter fonts",
    "✦ Glassmorphism card components",
    "✦ Gradient buttons with glow hover",
    "✦ Custom chat bubble styling",
    "✦ Dashboard metric cards",
    "✦ Progress bar system indicators",
    "✦ Responsive wide layout",
    "✦ Custom CSS injection (200+ lines)",
]
for item in ui_items:
    add_paragraph(tf9a, item, font_size=14, color=TEXT_MUTED)

add_card(slide9, 6.8, 1.7, 5.5, 5, BG_CARD)
tf9b = add_textbox(slide9, 7.1, 1.8, 5, 0.4, "FEATURES", font_size=14, bold=True, color=TEXT_GREEN)
feat_items = [
    "✦ Sidebar with API Key input",
    "✦ Real-time system status dashboard",
    "✦ Chat-based interaction interface",
    "✦ Styled AI response rendering",
    "✦ Error handling with user feedback",
    "✦ Session state management",
    "✦ @st.cache_data for performance",
    "✦ Hot-reload development",
    "✦ Mobile responsive layout",
]
for item in feat_items:
    add_paragraph(tf9b, item, font_size=14, color=TEXT_MUTED)

# ============================================================
# SLIDE 10: EXPECTED OUTPUT
# ============================================================
slide10 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide10)
add_accent_line(slide10, 0.8, 0.6, 2)
add_textbox(slide10, 0.8, 0.7, 8, 0.7, "EXPECTED OUTPUT & DEMO", font_size=32, bold=True, color=TEXT_WHITE)

add_card(slide10, 0.8, 1.7, 11.5, 5.2, BG_CARD)
tf10 = add_textbox(slide10, 1.1, 1.8, 11, 0.4, "SAMPLE OUTPUT: 3-DAY TRIP TO GOA (FEB 12-14)", font_size=14, bold=True, color=ACCENT_BLUE)

output_lines = [
    ("Flight Selected:", True, TEXT_GREEN),
    ("  IndiGo (₹4800) — Departs Delhi at 14:00", False, TEXT_MUTED),
    ("", False, TEXT_MUTED),
    ("Hotel Booked:", True, TEXT_GREEN),
    ("  Sea View Resort (₹3200/night, 4-star)", False, TEXT_MUTED),
    ("", False, TEXT_MUTED),
    ("Weather:", True, TEXT_GREEN),
    ("  Day 1: Sunny (31°C)  |  Day 2: Partly Cloudy  |  Day 3: Light Breeze", False, TEXT_MUTED),
    ("", False, TEXT_MUTED),
    ("Itinerary:", True, TEXT_GREEN),
    ("  Day 1: Baga Beach, Candolim Market", False, TEXT_MUTED),
    ("  Day 2: Basilica of Bom Jesus, Old Goa Heritage Walk", False, TEXT_MUTED),
    ("  Day 3: Water Sports at Calangute", False, TEXT_MUTED),
    ("", False, TEXT_MUTED),
    ("Budget Breakdown:", True, TEXT_GREEN),
    ("  Flight: ₹4,800  |  Hotel: ₹6,400  |  Food & Travel: ₹2,500", False, TEXT_MUTED),
    ("  ─────────────────────────────────────────────", False, BORDER_COLOR),
    ("  TOTAL COST: ₹13,700", True, ACCENT_BLUE),
]

for text, bold, color in output_lines:
    add_paragraph(tf10, text, font_size=14, bold=bold, color=color)

# ============================================================
# SLIDE 11: CODE QUALITY
# ============================================================
slide11 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide11)
add_accent_line(slide11, 0.8, 0.6, 2)
add_textbox(slide11, 0.8, 0.7, 8, 0.7, "CODE QUALITY & BEST PRACTICES", font_size=32, bold=True, color=TEXT_WHITE)

categories = [
    ("CODING STANDARDS", [
        "PEP 8 compliant formatting",
        "Meaningful variable & function names",
        "Modular code (functions & classes)",
        "try-except error handling",
        "Docstrings & inline comments",
    ], ACCENT_BLUE),
    ("DATABASE PRACTICES", [
        "Normalized tables (3NF)",
        "Indexed columns for fast queries",
        "Consistent naming conventions",
        "INSERT OR REPLACE for idempotency",
        "Parameterized queries (SQL injection safe)",
    ], ACCENT_PURPLE),
    ("PROJECT STRUCTURE", [
        "data/ — JSON datasets",
        "database.py — DB initialization",
        "tools.py — LangChain tool definitions",
        "agent.py — Agent configuration",
        "app.py — Streamlit frontend",
    ], TEXT_GREEN),
]

for i, (title, items, color) in enumerate(categories):
    x = 0.8 + i * 4.1
    add_card(slide11, x, 1.7, 3.7, 5, BG_CARD)
    add_textbox(slide11, x + 0.2, 1.8, 3.3, 0.4, title, font_size=14, bold=True, color=color, alignment=PP_ALIGN.CENTER)
    tf = add_textbox(slide11, x + 0.3, 2.4, 3.1, 0.3, "", font_size=13)
    for item in items:
        add_paragraph(tf, f"✦  {item}", font_size=13, color=TEXT_MUTED)

# ============================================================
# SLIDE 12: CONCLUSION & FUTURE SCOPE
# ============================================================
slide12 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide12)
add_accent_line(slide12, 0.8, 0.6, 2)
add_textbox(slide12, 0.8, 0.7, 8, 0.7, "CONCLUSION & FUTURE SCOPE", font_size=32, bold=True, color=TEXT_WHITE)

add_card(slide12, 0.8, 1.7, 5.5, 5, BG_CARD)
tf12a = add_textbox(slide12, 1.1, 1.8, 5, 0.4, "WHAT WE BUILT", font_size=14, bold=True, color=TEXT_GREEN)
conclusion = [
    "✦ Fully autonomous AI travel agent",
    "✦ 5 integrated LangChain tools",
    "✦ Real-time weather integration",
    "✦ SQLite-backed data layer",
    "✦ Premium dark-themed Streamlit UI",
    "✦ Structured itinerary generation",
    "✦ Budget estimation & optimization",
    "✦ Clean, modular, documented code",
]
for item in conclusion:
    add_paragraph(tf12a, item, font_size=14, color=TEXT_MUTED)

add_card(slide12, 6.8, 1.7, 5.5, 5, BG_CARD)
tf12b = add_textbox(slide12, 7.1, 1.8, 5, 0.4, "FUTURE ENHANCEMENTS", font_size=14, bold=True, color=ACCENT_BLUE)
future = [
    "✦ Multi-city trip planning",
    "✦ Real-time flight/hotel APIs",
    "✦ User authentication & profiles",
    "✦ Trip history & saved itineraries",
    "✦ Multi-language support",
    "✦ Voice-based trip planning",
    "✦ Integration with booking platforms",
    "✦ Advanced analytics dashboard",
]
for item in future:
    add_paragraph(tf12b, item, font_size=14, color=TEXT_MUTED)

# ============================================================
# SLIDE 13: THANK YOU
# ============================================================
slide13 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide13)
add_accent_line(slide13, 0, 0, 13.333)

add_textbox(slide13, 1, 2.0, 11, 1.2, "THANK YOU", font_size=56, bold=True, color=TEXT_WHITE, alignment=PP_ALIGN.CENTER, font_name="Calibri")
add_textbox(slide13, 1, 3.3, 11, 0.7, "Agentic AI-Based Travel Planning Assistant", font_size=22, color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER)

add_card(slide13, 3.5, 4.5, 6.3, 1.8, BG_CARD)
tf13 = add_textbox(slide13, 3.7, 4.6, 5.9, 0.4, "PROJECT LINKS", font_size=14, bold=True, color=ACCENT_PURPLE, alignment=PP_ALIGN.CENTER)
add_paragraph(tf13, "", font_size=6)
add_paragraph(tf13, "GitHub: github.com/viRAJ357/AIPLANNER", font_size=15, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)
add_paragraph(tf13, "Tech: Python | LangChain | Streamlit | SQLite", font_size=15, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

add_textbox(slide13, 1, 6.8, 11, 0.4, "Questions & Discussion", font_size=16, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# ============================================================
# SAVE
# ============================================================
prs.save("AI_Travel_Planner_Report.pptx")
print("PPT generated successfully: AI_Travel_Planner_Report.pptx")
